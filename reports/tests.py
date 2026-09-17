"""Тесты статистики, выгрузки и карты."""

from django.contrib.auth import get_user_model
from django.test import TestCase

from directory.models import Category, District
from tickets.models import Channel, Citizen, Message, Ticket

User = get_user_model()


class ReportsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(email="op@meriya.kg",
                                            password="pass12345")
        cat = Category.objects.create(name="Дороги", slug="roads")
        dist = District.objects.create(name="Центр", slug="center")
        citizen = Citizen.objects.create(tg_user_id=1, chat_id=1, first_name="А")
        cls.ticket = Ticket.objects.create(
            citizen=citizen, title="Яма на Ленина", category=cat,
            district=dist, address="Ленина 5")
        message = Message.objects.create(ticket=cls.ticket, author="ai",
                                         text="ответ", rating="up")

    def setUp(self):
        self.client.login(email="op@meriya.kg", password="pass12345")


class StatsTests(ReportsTestCase):
    def test_page_renders_with_data(self):
        response = self.client.get("/stats/")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("Дороги", html)
        self.assertIn("Центр", html)
        self.assertIn("100%", html)  # доля «помогло»

    def test_requires_login(self):
        self.client.logout()
        self.assertEqual(self.client.get("/stats/").status_code, 302)


class ExportTests(ReportsTestCase):
    def test_xlsx_download(self):
        response = self.client.get("/export/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(response.content)).active
        rows = list(ws.values)
        self.assertEqual(rows[0][0], "Номер")
        self.assertEqual(rows[0][1], "Канал")
        self.assertEqual(rows[1][1], "Telegram")
        self.assertEqual(rows[1][3], "Яма на Ленина")

    def test_export_respects_filters(self):
        response = self.client.get("/export/?q=несуществующее")
        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(response.content)).active
        self.assertEqual(len(list(ws.values)), 1, "только заголовок")

    def test_export_filters_by_channel(self):
        wa_citizen = Citizen.objects.create(channel=Channel.WHATSAPP, chat_id=996700123456)
        Ticket.objects.create(citizen=wa_citizen, title="WhatsApp-обращение",
                              channel=Channel.WHATSAPP)
        from io import BytesIO
        from openpyxl import load_workbook

        response = self.client.get("/export/?channel=whatsapp")
        ws = load_workbook(BytesIO(response.content)).active
        rows = list(ws.values)
        self.assertEqual(len(rows), 2)  # заголовок + 1 WhatsApp-заявка
        self.assertEqual(rows[1][3], "WhatsApp-обращение")


class MapTests(ReportsTestCase):
    def test_empty_state(self):
        response = self.client.get("/map/")
        self.assertContains(response, "нет координат")

    def test_points_rendered(self):
        Ticket.objects.filter(pk=self.ticket.pk).update(lat=40.93, lon=73.0)
        response = self.client.get("/map/")
        self.assertContains(response, "40.93")
        self.assertContains(response, self.ticket.number)

    def test_failed_geocode_hidden(self):
        Ticket.objects.filter(pk=self.ticket.pk).update(lat=-1000, lon=-1000)
        response = self.client.get("/map/")
        self.assertContains(response, "нет координат")


class GeocodeBoundsTests(TestCase):
    """
    Проверка рамки города. Без неё геокодер находил «ул. Ленина» за сотню
    километров от Манаса, и заявка вставала на карте в чужом районе.
    """

    def test_city_points_accepted(self):
        from reports.management.commands.geocode import inside_city
        self.assertTrue(inside_city(40.9333, 72.9833))   # центр города
        self.assertTrue(inside_city(40.92368, 73.00326))  # ул. Токтогула

    def test_far_points_rejected(self):
        from reports.management.commands.geocode import inside_city
        self.assertFalse(inside_city(41.34717, 72.22169),  # была такая ошибка
                         "точка за сотню километров не должна попадать на карту")
        self.assertFalse(inside_city(42.87, 74.59))        # Бишкек
        self.assertFalse(inside_city(0, 0))                # пустой ответ сервиса

    def test_bounds_are_sane(self):
        from reports.management.commands.geocode import (
            CITY_LAT, CITY_LON, LAT_MAX, LAT_MIN, LON_MAX, LON_MIN)
        self.assertTrue(LAT_MIN < CITY_LAT < LAT_MAX)
        self.assertTrue(LON_MIN < CITY_LON < LON_MAX)
        self.assertLess(LAT_MAX - LAT_MIN, 1.0, "рамка не должна быть на пол-страны")


class TemplateHygieneTests(TestCase):
    """
    Страницы не должны показывать служебный текст.

    Появилось после реального случая: многострочный комментарий {# ... #}
    Django комментарием НЕ считает (только однострочный) и вывел пояснение
    для разработчика прямо в таблицу обращений на глазах у заказчика.
    Для многострочных есть тег comment.
    """

    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(
            email="adm@meriya.kg", password="pass12345", role="admin")

    def setUp(self):
        self.client.login(email="adm@meriya.kg", password="pass12345")

    def test_no_leaked_comment_markers_on_pages(self):
        for url in ("/", "/tickets/new/", "/bot/settings/", "/bot/keys/",
                    "/bot/answers/", "/broadcasts/", "/stats/", "/map/",
                    "/directory/", "/staff/", "/audit/"):
            with self.subTest(url=url):
                html = self.client.get(url).content.decode()
                for marker in ("{#", "#}", "{%", "%}"):
                    self.assertNotIn(marker, html,
                                     f"на странице {url} видна разметка шаблона")

    def test_templates_have_no_multiline_hash_comments(self):
        """Однострочный {# … #} допустим, многострочный — нет."""
        from django.conf import settings
        offenders = []
        for template_dir in [settings.BASE_DIR / "templates"]:
            for path in template_dir.rglob("*.html"):
                for number, line in enumerate(path.read_text().splitlines(), 1):
                    if "{#" in line and "#}" not in line:
                        offenders.append(f"{path.name}:{number}")
        self.assertEqual(offenders, [],
                         "многострочные {# #} выводятся на страницу как текст")
