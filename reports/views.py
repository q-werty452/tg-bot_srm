"""
reports/views.py — статистика, карта и выгрузка в Excel.

Графики статистики — обычные полоски на CSS, без внешних библиотек:
панель обязана работать без интернета. Единственное исключение — карта:
ей нужны тайлы OpenStreetMap из сети, о чём страница честно предупреждает.
"""

import json
from datetime import timedelta

from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from directory.models import Category, District
from tickets.models import Channel, Message, Ticket


def _bars(rows, label_key, count_key="n"):
    """Подготовить строки для CSS-графика: подписи + ширина в процентах."""
    top = max((r[count_key] for r in rows), default=0) or 1
    return [
        {"label": r[label_key] or "—", "count": r[count_key],
         "percent": round(r[count_key] * 100 / top)}
        for r in rows
    ]


@login_required
def stats(request):
    now = timezone.now()
    month_ago = now - timedelta(days=30)
    tickets = Ticket.objects.filter(created_at__gte=month_ago)

    # Плитки: где мы сейчас.
    tiles = {
        "total": tickets.count(),
        "open": Ticket.objects.open().count(),
        "overdue": Ticket.objects.overdue().count(),
        "done": tickets.filter(status=Ticket.Status.DONE).count(),
    }

    # Оценки жителей под ответами ИИ.
    ratings = Message.objects.filter(
        author=Message.Author.AI, created_at__gte=month_ago, rating__isnull=False)
    up = ratings.filter(rating="up").count()
    total_rated = ratings.count()
    tiles["rating"] = f"{round(up * 100 / total_rated)}%" if total_rated else "—"
    tiles["rated_count"] = total_rated

    # Обращения по дням за две недели.
    by_day = []
    for shift in range(13, -1, -1):
        day = (now - timedelta(days=shift)).date()
        n = Ticket.objects.filter(created_at__date=day).count()
        by_day.append({"label": day.strftime("%d.%m"), "n": n})
    top = max((d["n"] for d in by_day), default=0) or 1
    for d in by_day:
        d["percent"] = round(d["n"] * 100 / top)

    # По категориям и районам (за месяц).
    by_category = _bars(
        list(tickets.values("category__name").annotate(n=Count("id"))
             .order_by("-n")[:10]), "category__name")
    by_district = _bars(
        list(tickets.values("district__name").annotate(n=Count("id"))
             .order_by("-n")[:10]), "district__name")

    # По часам суток: когда жителям удобно писать. Считаем в Python —
    # это переживёт переезд с SQLite на PostgreSQL без правок.
    hour_counts: dict[int, int] = {}
    for created in tickets.values_list("created_at", flat=True):
        hour = timezone.localtime(created).hour
        hour_counts[hour] = hour_counts.get(hour, 0) + 1
    top_h = max(hour_counts.values(), default=0) or 1
    by_hour = [{"label": f"{h:02d}", "n": hour_counts.get(h, 0),
                "percent": round(hour_counts.get(h, 0) * 100 / top_h)}
               for h in range(24)]

    return render(request, "stats.html", {
        "section": "stats", "tiles": tiles, "by_day": by_day,
        "by_category": by_category, "by_district": by_district,
        "by_hour": by_hour,
    })


def _filtered_tickets(request):
    """Те же фильтры, что на дашборде, — выгрузка отдаёт то, что видно."""
    qs = Ticket.objects.select_related(
        "citizen", "category", "district", "executor", "assignee")
    status = request.GET.get("status", "")
    if status == "overdue":
        qs = qs.overdue()
    elif status in Ticket.Status.values:
        qs = qs.filter(status=status)
    if request.GET.get("category", "").isdigit():
        qs = qs.filter(category_id=request.GET["category"])
    if request.GET.get("district", "").isdigit():
        qs = qs.filter(district_id=request.GET["district"])
    channel = request.GET.get("channel", "")
    if channel in Channel.values:
        qs = qs.filter(channel=channel)
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(Q(number__icontains=q) | Q(title__icontains=q)
                       | Q(description__icontains=q) | Q(address__icontains=q)
                       | Q(citizen__first_name__icontains=q)
                       | Q(citizen__last_name__icontains=q)
                       | Q(citizen__phone__icontains=q))
    return qs.order_by("-created_at")


@login_required
def export_xlsx(request):
    """Выгрузка обращений в Excel — для отчётности мэрии."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Обращения"
    headers = ["Номер", "Канал", "Создана", "Заголовок", "Категория", "Район", "Адрес",
               "Статус", "Просрочена", "Исполнитель", "Ответственный",
               "Житель", "Телефон", "Кто отвечает"]
    ws.append(headers)

    for t in _filtered_tickets(request):
        ws.append([
            t.number,
            t.get_channel_display(),
            timezone.localtime(t.created_at).strftime("%d.%m.%Y %H:%M"),
            t.title,
            t.category.name if t.category else "",
            t.district.name if t.district else "",
            t.address,
            t.get_status_display(),
            "да" if t.is_overdue else "",
            str(t.executor) if t.executor else "",
            str(t.assignee) if t.assignee else "",
            str(t.citizen),
            t.citizen.phone,
            t.get_answer_mode_display(),
        ])

    # Ширина колонок по содержимому, без фанатизма.
    for col, header in enumerate(headers, start=1):
        width = max(len(header), *(len(str(c.value or "")) for c in
                                   ws[get_column_letter(col)])) + 2
        ws.column_dimensions[get_column_letter(col)].width = min(width, 45)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument"
                     ".spreadsheetml.sheet")
    filename = f"obrashcheniya_{timezone.localdate().isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def tickets_map(request):
    """Карта обращений. Точки — заявки, у которых определены координаты."""
    points = [
        {"lat": t.lat, "lon": t.lon, "number": t.number, "title": t.title,
         "status": t.get_status_display(), "overdue": t.is_overdue}
        for t in Ticket.objects.filter(lat__isnull=False, lon__isnull=False)
                                .exclude(lat__lt=-90)  # -1000 = «адрес не распознан»
                                .select_related("category")
    ]
    without = Ticket.objects.filter(lat__isnull=True).exclude(address="").count()
    return render(request, "map.html", {
        "section": "map", "points": points,
        "points_json": json.dumps(points, ensure_ascii=False),
        "without_coords": without,
    })
